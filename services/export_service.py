"""
Exportacion de los datos de un despacho a un libro de Excel (.xlsx),
para que el especialista/liquidador pueda sacar la informacion fuera de
la app (ninguna otra forma de hacerlo existia antes de esto).

La exportacion a JSON no pasa por aqui: el frontend ya tiene el
`DespachoDetalleOut` completo cargado (misma llamada que arma este
mismo modulo, ver `app/main.py::_armar_detalle_despacho`), asi que arma
el archivo JSON del lado del cliente sin ida y vuelta al backend. Un
.xlsx si necesita una libreria del lado del servidor (openpyxl).
"""
from __future__ import annotations

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

_FUENTE_ENCABEZADO = Font(bold=True)


def _a_dict(valor: object) -> dict:
    """`despacho`/`documentos`/`borrador`/`decision` en `detalle` ya son
    dicts crudos de Supabase, pero `validaciones` (list[ResultadoValidacion])
    y `clasificacion` (PropuestaClasificacion | None) son objetos Pydantic
    -- se usa la misma funcion para normalizar cualquiera de los dos casos
    antes de leer sus campos."""
    if isinstance(valor, BaseModel):
        return valor.model_dump()
    return valor


def _escribir_par(hoja: Worksheet, fila: int, etiqueta: str, valor: object) -> int:
    """Escribe una fila `etiqueta | valor` y devuelve la siguiente fila
    libre. Valores que no son texto/numero plano (listas, dicts) se
    serializan a JSON en la misma celda -- generico para cualquier schema
    de documento sin tener que mapear columna por columna."""
    hoja.cell(row=fila, column=1, value=etiqueta).font = _FUENTE_ENCABEZADO
    if isinstance(valor, (dict, list)):
        valor = json.dumps(valor, ensure_ascii=False)
    hoja.cell(row=fila, column=2, value=valor)
    return fila + 1


def _hoja_despacho(libro: Workbook, despacho: dict) -> None:
    hoja = libro.create_sheet("Despacho")
    fila = 1
    for etiqueta, clave in (
        ("Número de despacho", "numero_despacho"),
        ("Cliente", "cliente"),
        ("Estado", "estado"),
        ("Fecha de creación", "fecha_creacion"),
    ):
        fila = _escribir_par(hoja, fila, etiqueta, despacho.get(clave))
    hoja.column_dimensions["A"].width = 22
    hoja.column_dimensions["B"].width = 50


def _hoja_documentos(libro: Workbook, documentos: list[dict]) -> None:
    hoja = libro.create_sheet("Documentos")
    fila = 1
    if not documentos:
        hoja.cell(row=1, column=1, value="Sin documentos cargados.")
        return

    for documento in documentos:
        hoja.cell(row=fila, column=1, value=documento["tipo_documento"]).font = Font(bold=True, size=13)
        fila += 1
        fila = _escribir_par(hoja, fila, "Procesado", documento.get("procesado"))
        fila = _escribir_par(hoja, fila, "Método de extracción", documento.get("metodo_extraccion"))
        for clave, valor in (documento.get("contenido_json") or {}).items():
            fila = _escribir_par(hoja, fila, clave, valor)
        fila += 1  # fila en blanco entre documentos
    hoja.column_dimensions["A"].width = 28
    hoja.column_dimensions["B"].width = 60


def _hoja_validaciones(libro: Workbook, validaciones: list[dict]) -> None:
    hoja = libro.create_sheet("Validaciones")
    hoja.append(["Regla", "Severidad", "Detalle"])
    for celda in hoja[1]:
        celda.font = _FUENTE_ENCABEZADO
    for v in validaciones:
        v = _a_dict(v)
        hoja.append([v.get("regla"), v.get("severidad"), v.get("detalle")])
    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 12
    hoja.column_dimensions["C"].width = 90


def _hoja_clasificacion(libro: Workbook, clasificacion: object | None, decision: dict | None) -> None:
    hoja = libro.create_sheet("Clasificación")
    fila = 1
    if clasificacion:
        clasificacion = _a_dict(clasificacion)
        fila = _escribir_par(hoja, fila, "Subpartida sugerida", clasificacion.get("subpartida_sugerida"))
        score = clasificacion.get("score_confianza")
        confianza_texto = clasificacion.get("nivel_confianza")
        if score is not None:
            confianza_texto = f"{confianza_texto} ({round(score * 100)}%)"
        fila = _escribir_par(hoja, fila, "Confianza", confianza_texto)
        fila = _escribir_par(hoja, fila, "Sustento legal (RGI)", clasificacion.get("sustento_legal_rgi"))
        fila = _escribir_par(
            hoja, fila, "Información faltante", clasificacion.get("informacion_faltante_alert") or []
        )
    elif decision:
        fila = _escribir_par(hoja, fila, "Subpartida final", decision.get("subpartida_final_humano"))
        fila = _escribir_par(hoja, fila, "Decisión", decision.get("tipo_accion"))
        fila = _escribir_par(hoja, fila, "Motivo de la observación", decision.get("motivo_modificacion"))
    else:
        hoja.cell(row=1, column=1, value="Este despacho todavía no tiene una propuesta de clasificación.")
    hoja.column_dimensions["A"].width = 24
    hoja.column_dimensions["B"].width = 70


def generar_excel_despacho(detalle: dict) -> bytes:
    """`detalle` es el mismo dict que arma `_armar_detalle_despacho` en
    app/main.py (despacho/documentos/validaciones/clasificacion/decision/
    borrador) -- se reusa tal cual, sin volver a golpear la base de datos."""
    libro = Workbook()
    libro.remove(libro.active)  # la hoja "Sheet" en blanco que trae por defecto

    _hoja_despacho(libro, detalle["despacho"])
    _hoja_documentos(libro, detalle["documentos"])
    _hoja_validaciones(libro, detalle["validaciones"])
    _hoja_clasificacion(libro, detalle.get("clasificacion"), detalle.get("decision"))

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
